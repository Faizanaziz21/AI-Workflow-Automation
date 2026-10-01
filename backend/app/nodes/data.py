"""Data nodes: HTTP request, CSV, JSON transformer, file reader, Excel."""

from __future__ import annotations

import base64
import csv
import io
import json
from typing import Any, Literal

from pydantic import BaseModel, ConfigDict, Field

from app.connectors.builtin.http_rest.client import perform_request
from app.connectors.builtin.http_rest.schemas import RequestInput
from app.connectors.sdk import ConnectorError
from app.core.egress import get_http_client
from app.engine.expressions import ExpressionError, evaluate, is_template, render
from app.nodes.base import Completed, NodeContext, NodeError, NodeResult, NodeType

MAX_ROWS = 100_000


class HttpRequestConfig(BaseModel):
    model_config = ConfigDict(extra="forbid")
    connection_id: str | None = Field(
        default=None,
        description="Optional REST API connection providing base URL + auth",
        json_schema_extra={"x-connection": "http_rest"},
    )
    method: Literal["GET", "POST", "PUT", "PATCH", "DELETE", "HEAD"] = "GET"
    url: str = Field(description="Absolute URL, or a path when a connection is selected")
    query: dict[str, Any] = Field(default_factory=dict)
    headers: dict[str, str] = Field(default_factory=dict)
    body: Any = Field(default=None, description="JSON body (object/array) or text")
    body_type: Literal["json", "form", "text", "none"] = "json"
    timeout_seconds: float = Field(default=30, gt=0, le=300)
    fail_on_http_error: bool = True
    response_format: Literal["auto", "json", "text"] = "auto"


class HttpRequestNode(NodeType):
    type = "data.http_request"
    category = "data"
    label = "HTTP request"
    description = "Call any HTTP API (egress-policy enforced; Idempotency-Key sent on writes)."
    icon = "globe"
    config_model = HttpRequestConfig
    output_schema = {
        "type": "object",
        "properties": {
            "status": {"type": "integer"},
            "ok": {"type": "boolean"},
            "headers": {"type": "object"},
            "body": {},
        },
    }

    async def execute(self, ctx: NodeContext, config: HttpRequestConfig) -> NodeResult:
        req = RequestInput(
            method=config.method,
            path=config.url,
            query=config.query,
            headers=config.headers,
            json_body=config.body if config.body_type == "json" and config.body is not None else None,
            form_body={k: str(v) for k, v in config.body.items()}
            if config.body_type == "form" and isinstance(config.body, dict)
            else None,
            text_body=str(config.body) if config.body_type == "text" and config.body is not None else None,
            timeout_seconds=config.timeout_seconds,
            fail_on_http_error=config.fail_on_http_error,
            response_format=config.response_format,
        )
        try:
            if config.connection_id:
                resolved = await ctx.runtime.connection(config.connection_id)
                if resolved.connector.key != "http_rest":
                    raise NodeError("HTTP request node requires a 'REST API' connection")
                resolved.context.idempotency_key = ctx.idempotency_key
                out = await resolved.connector.run_action("request", resolved.context, req.model_dump())
                return Completed(output=out)
            if not config.url.startswith(("http://", "https://")):
                raise NodeError("URL must be absolute when no connection is selected")
            result = await perform_request(get_http_client(), config.url, req, idempotency_key=ctx.idempotency_key)
        except ConnectorError as exc:
            raise NodeError(exc.message, retryable=exc.retryable, details=exc.details) from exc
        return Completed(output=result.model_dump(mode="json"))


class CsvConfig(BaseModel):
    model_config = ConfigDict(extra="forbid")
    operation: Literal["parse", "build"] = "parse"
    text: str | None = Field(default=None, description="CSV text (parse)")
    file_id: str | None = Field(default=None, description="Stored file id (parse)")
    rows: list[dict[str, Any]] | None = Field(default=None, description="Rows to serialise (build)")
    delimiter: str = Field(default=",", min_length=1, max_length=1)
    has_header: bool = True
    max_rows: int = Field(default=10_000, ge=1, le=MAX_ROWS)
    save_as_file: str | None = Field(default=None, description="Build: store result as a file with this name")


class CsvNode(NodeType):
    type = "data.csv"
    category = "data"
    label = "CSV parser"
    description = "Parse CSV text/files into rows, or build CSV from rows."
    icon = "table"
    config_model = CsvConfig
    side_effects = False
    output_schema = {
        "type": "object",
        "properties": {
            "rows": {"type": "array"},
            "count": {"type": "integer"},
            "columns": {"type": "array"},
            "csv": {"type": "string"},
        },
    }

    async def execute(self, ctx: NodeContext, config: CsvConfig) -> NodeResult:
        if config.operation == "build":
            rows = config.rows or []
            columns: list[str] = []
            for r in rows:
                for k in r:
                    if k not in columns:
                        columns.append(k)
            buf = io.StringIO()
            writer = csv.DictWriter(buf, fieldnames=columns, delimiter=config.delimiter, extrasaction="ignore")
            writer.writeheader()
            for r in rows:
                writer.writerow({k: json.dumps(v) if isinstance(v, (dict, list)) else v for k, v in r.items()})
            out: dict[str, Any] = {"csv": buf.getvalue(), "count": len(rows), "columns": columns}
            if config.save_as_file:
                out["file"] = await ctx.runtime.files.save(
                    config.save_as_file, buf.getvalue().encode(), "text/csv", "workflow"
                )
                out["csv"] = out["csv"][:10_000]
            return Completed(output=out)
        if config.text is None and not config.file_id:
            raise NodeError("Provide CSV text or file_id")
        text = config.text
        if text is None:
            _, data = await ctx.runtime.files.load(str(config.file_id))
            text = data.decode("utf-8-sig", errors="replace")
        reader = csv.reader(io.StringIO(text), delimiter=config.delimiter)
        records = list(_take(reader, config.max_rows + 1))
        if not records:
            return Completed(output={"rows": [], "count": 0, "columns": []})
        if config.has_header:
            header = [h.strip() for h in records[0]]
            body = records[1:]
            rows = [dict(zip(header, r + [""] * (len(header) - len(r)), strict=False)) for r in body]
        else:
            header = [f"col_{i + 1}" for i in range(max(len(r) for r in records))]
            rows = [dict(zip(header, r, strict=False)) for r in records]
        truncated = len(rows) > config.max_rows
        rows = rows[: config.max_rows]
        return Completed(output={"rows": rows, "count": len(rows), "columns": header, "truncated": truncated})


def _take(it: Any, n: int) -> Any:
    for i, x in enumerate(it):
        if i >= n:
            return
        yield x


class JsonTransformConfig(BaseModel):
    model_config = ConfigDict(extra="forbid")
    mode: Literal["object", "map_list"] = "object"
    template: Any = Field(description="Output shape; any value may contain expressions")
    items: Any = Field(default=None, description="map_list: list to map; template evaluated per `item`")


class JsonTransformNode(NodeType):
    type = "data.json_transform"
    category = "data"
    label = "JSON transformer"
    description = "Reshape data with a template of expressions (object or per-item mapping)."
    icon = "braces"
    config_model = JsonTransformConfig
    raw_fields = frozenset({"template"})
    side_effects = False

    async def execute(self, ctx: NodeContext, config: JsonTransformConfig) -> NodeResult:
        try:
            if config.mode == "object":
                return Completed(output=render(config.template, ctx.data))
            items = config.items or []
            if not isinstance(items, list):
                raise NodeError("map_list mode requires 'items' to be a list")
            return Completed(
                output={
                    "items": [
                        render(config.template, {**ctx.data, "item": item, "index": i}) for i, item in enumerate(items)
                    ]
                }
            )
        except ExpressionError as exc:
            raise NodeError(f"Transform failed: {exc.message}") from exc


class FileReadConfig(BaseModel):
    model_config = ConfigDict(extra="forbid")
    file_id: str
    read_as: Literal["auto", "text", "json", "base64"] = "auto"
    max_chars: int = Field(default=200_000, ge=1, le=2_000_000)


def extract_text(content: bytes, content_type: str, filename: str) -> str:
    name = filename.lower()
    if content_type == "application/pdf" or name.endswith(".pdf"):
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(content))
            return "\n\n".join((page.extract_text() or "") for page in reader.pages[:200])
        except Exception as exc:
            raise NodeError(f"Could not read PDF: {exc}") from exc
    if name.endswith((".xlsx", ".xlsm")):
        rows = read_excel(content, None, True, 2000)
        return "\n".join(json.dumps(r, default=str) for r in rows)
    return content.decode("utf-8-sig", errors="replace")


class FileReadNode(NodeType):
    type = "data.file_read"
    category = "data"
    label = "File reader"
    description = "Read a stored file as text (PDF/Excel text extraction), JSON or base64."
    icon = "file-text"
    config_model = FileReadConfig
    side_effects = False
    output_schema = {
        "type": "object",
        "properties": {
            "file": {"type": "object"},
            "text": {"type": "string"},
            "json": {},
            "base64": {"type": "string"},
        },
    }

    async def execute(self, ctx: NodeContext, config: FileReadConfig) -> NodeResult:
        meta, content = await ctx.runtime.files.load(config.file_id)
        out: dict[str, Any] = {"file": meta}
        mode = config.read_as
        if mode == "auto":
            ctype = meta.get("content_type", "")
            mode = (
                "json"
                if "json" in ctype
                else ("base64" if ctype.startswith(("image/", "audio/", "video/")) else "text")
            )
        if mode == "base64":
            out["base64"] = base64.b64encode(content).decode()[: config.max_chars]
        elif mode == "json":
            try:
                out["json"] = json.loads(content)
            except json.JSONDecodeError as exc:
                raise NodeError(f"File is not valid JSON: {exc.msg}") from exc
        else:
            text = extract_text(content, meta.get("content_type", ""), meta.get("filename", ""))
            out["text"] = text[: config.max_chars]
            out["truncated"] = len(text) > config.max_chars
        return Completed(output=out)


def read_excel(content: bytes, sheet: str | None, header: bool, max_rows: int) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    out: list[dict[str, Any]] = []
    columns: list[str] | None = None
    for i, row in enumerate(rows_iter):
        if header and columns is None:
            columns = [str(c) if c is not None else f"col_{j + 1}" for j, c in enumerate(row)]
            continue
        cols = columns or [f"col_{j + 1}" for j in range(len(row))]
        out.append(
            {
                cols[j] if j < len(cols) else f"col_{j + 1}": (v.isoformat() if hasattr(v, "isoformat") else v)
                for j, v in enumerate(row)
            }
        )
        if len(out) >= max_rows or i > MAX_ROWS:
            break
    wb.close()
    return out


class ExcelConfig(BaseModel):
    model_config = ConfigDict(extra="forbid")
    operation: Literal["read", "write"] = "read"
    file_id: str | None = None
    sheet: str | None = None
    has_header: bool = True
    max_rows: int = Field(default=10_000, ge=1, le=MAX_ROWS)
    rows: list[dict[str, Any]] | None = None
    filename: str = "export.xlsx"


class ExcelNode(NodeType):
    type = "data.excel"
    category = "data"
    label = "Excel"
    description = "Read rows from an .xlsx workbook or write rows to a new workbook file."
    icon = "sheet"
    config_model = ExcelConfig
    side_effects = False

    async def execute(self, ctx: NodeContext, config: ExcelConfig) -> NodeResult:
        if config.operation == "read":
            if not config.file_id:
                raise NodeError("file_id is required to read a workbook")
            _, content = await ctx.runtime.files.load(config.file_id)
            try:
                rows = read_excel(content, config.sheet, config.has_header, config.max_rows)
            except KeyError as exc:
                raise NodeError(f"Sheet not found: {config.sheet}") from exc
            except Exception as exc:
                raise NodeError(f"Could not read workbook: {exc}") from exc
            return Completed(output={"rows": rows, "count": len(rows)})
        from openpyxl import Workbook

        rows = config.rows or []
        columns: list[str] = []
        for r in rows:
            columns.extend(k for k in r if k not in columns)
        wb = Workbook()
        ws = wb.active
        ws.title = (config.sheet or "Sheet1")[:31]
        ws.append(columns)
        for r in rows:
            ws.append([json.dumps(v) if isinstance(v, (dict, list)) else v for v in (r.get(c) for c in columns)])
        buf = io.BytesIO()
        wb.save(buf)
        ref = await ctx.runtime.files.save(
            config.filename,
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "workflow",
        )
        return Completed(output={"file": ref, "count": len(rows)})


DATA_NODES: list[type[NodeType]] = [HttpRequestNode, CsvNode, JsonTransformNode, FileReadNode, ExcelNode]

__all__ = ["DATA_NODES", "evaluate", "extract_text", "is_template"]
